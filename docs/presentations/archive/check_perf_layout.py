"""Structural layout check of both VSX3000 performance-testing decks, for use when the slides cannot be rendered.

For every slide it reports
  - two text boxes whose boxes overlap (a text box is any shape holding text; a figure counts as one);
  - a text box whose estimated text height exceeds the box height;
  - a table cell whose estimated text height exceeds its row height, and a table that ends below the content area;
  - a shape that reaches beyond the slide's margins.

The line count uses the deck generator's own text-width model (make_perf_deck.js, countLines):
characters per line = box width / (font size x CHAR_W), words wrapped greedily. The model is deliberately a
little wide, so a warning means "might not fit"; a render is the final judge.

Run from anywhere:  python3 docs/presentations/archive/check_perf_layout.py
"""
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parents[3]  # archive/ -> presentations/ -> docs/ -> repository root
DECKS = [
    REPO / "docs/presentations/vsx3000_procurement_build.pptx",
    REPO / "docs/presentations/vsx3000_test_procedure.pptx",
]
# The generator's constants (make_perf_deck.js)
CHAR_W_REGULAR = 0.46  # average character width as a fraction of the font size, regular text
CHAR_W_BOLD = 0.5  # the same, bold text
BODY_LINE_FACTOR = 1.2  # line height as a multiple of the font size
DEFAULT_PT = 16  # BODY_PT, the generator's default text size
BULLET_INDENT_PT = 14  # hanging indent of bullets
SLIDE_H = 7.5
SLIDE_W = 13.333
EDGE_MARGIN = 0.5 - 0.01  # content must stay 0.5 in from the slide's edges (a hair of tolerance); the slide number sits below
CONTENT_BOTTOM = 6.7  # last y available to content (slide height - number box - gaps)
OVERLAP_TOLERANCE = 0.02  # inches of overlap ignored (touching boxes)
FIT_TOLERANCE = 0.02  # inches of overflow ignored
IN = 914400.0  # EMU per inch


def count_lines(text, width_in, pt, bold):
    """Same greedy wrap as the generator's countLines."""
    char_w = pt * (CHAR_W_BOLD if bold else CHAR_W_REGULAR) / 72
    max_chars = max(1, math.floor(width_in / char_w))
    lines, used = 1, 0
    for word in text.split(" "):
        need = len(word) if used == 0 else used + 1 + len(word)
        if need <= max_chars:
            used = need
        else:
            lines += 0 if used == 0 else 1
            used = len(word)
    return lines


def paragraph_height(par, width_in, last):
    """Estimated height (in) of one paragraph: its lines at its largest font size, plus the space after it."""
    runs = [r for r in par.runs if r.text]
    text = "".join(r.text for r in runs).replace(" ", " ")
    if not text:
        return 0.0
    pt = max((r.font.size.pt for r in runs if r.font.size), default=DEFAULT_PT)
    bold = all(bool(r.font.bold) for r in runs)
    pPr = par._p.pPr
    indent_in = 0.0
    space_pt = 0.0
    if pPr is not None:
        if pPr.get("marL"):
            indent_in = int(pPr.get("marL")) / IN
        spc = pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcAft")
        if spc is not None:
            pts = spc.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcPts")
            if pts is not None:
                space_pt = int(pts.get("val")) / 100
    lines = count_lines(text, width_in - indent_in, pt, bold)
    return lines * pt * BODY_LINE_FACTOR / 72 + (0 if last else space_pt / 72)


def frame_height(tf, width_in):
    pars = list(tf.paragraphs)
    return sum(paragraph_height(p, width_in, i == len(pars) - 1) for i, p in enumerate(pars))


def box(shape):
    return shape.left / IN, shape.top / IN, shape.width / IN, shape.height / IN


def overlap(a, b):
    ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return (ox, oy) if ox > OVERLAP_TOLERANCE and oy > OVERLAP_TOLERANCE else None


def check_slide(slide, number):
    warnings = []
    texts = []  # (name, box) of every shape that holds text
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            x, y, w, h = box(sh)
            texts.append((sh.name, (x, y, w, h)))
            need = frame_height(sh.text_frame, w)
            if need > h + FIT_TOLERANCE:
                warnings.append(f"text may overflow: {sh.name!r} needs {need:.2f} in, box is {h:.2f} in")
        if sh.shape_type == 13 and not sh.name.endswith(("glyph", "check")):  # a figure (picture), not an icon glyph: counts as a box too
            texts.append((sh.name, box(sh)))
        if getattr(sh, "has_table", False) and sh.has_table:
            tbl = sh.table
            x, y, w, h = box(sh)
            texts.append((sh.name, (x, y, w, sum(r.height for r in tbl.rows) / IN)))
            total = sum(r.height for r in tbl.rows) / IN
            if y + total > CONTENT_BOTTOM + FIT_TOLERANCE:
                warnings.append(f"table {sh.name!r} ends at {y + total:.2f} in, below the content area ({CONTENT_BOTTOM} in)")
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row.cells):
                    cw = tbl.columns[ci].width / IN - (cell.margin_left + cell.margin_right) / IN
                    need = frame_height(cell.text_frame, cw) + (cell.margin_top + cell.margin_bottom) / IN
                    if need > row.height / IN + FIT_TOLERANCE:
                        warnings.append(f"table cell may overflow: {sh.name!r} row {ri + 1} col {ci + 1} needs {need:.2f} in, row is {row.height / IN:.2f} in")
        x, y, w, h = box(sh)
        if sh.name.startswith("Slide Number"):
            continue  # the slide number sits in the bottom margin by design
        if x < EDGE_MARGIN or y < EDGE_MARGIN or x + w > SLIDE_W - EDGE_MARGIN or y + h > SLIDE_H - EDGE_MARGIN:
            warnings.append(f"shape {sh.name!r} reaches into the slide margin: x {x:.2f}-{x + w:.2f}, y {y:.2f}-{y + h:.2f}")
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            o = overlap(texts[i][1], texts[j][1])
            if o:
                warnings.append(f"text boxes overlap: {texts[i][0]!r} and {texts[j][0]!r} by {o[0]:.2f} x {o[1]:.2f} in")
    return warnings, len(texts)


def main():
    total = 0
    for path in DECKS:
        prs = Presentation(str(path))
        print(f"{path.name}: {len(prs.slides)} slides")
        for n, slide in enumerate(prs.slides, 1):
            warnings, n_text = check_slide(slide, n)
            total += len(warnings)
            title = next((sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")  # the title is the first text shape
            print(f"  slide {n:2d} ({n_text} text boxes) {title[:60]!r}: {'ok' if not warnings else str(len(warnings)) + ' warning(s)'}")
            for w in warnings:
                print("      WARNING: " + w)
    print(f"total warnings: {total}")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
